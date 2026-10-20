#!/usr/bin/env python3
"""EXPOSED source-native encoding calibration after V1 UNKNOWN; not a prospective test."""
from collections import Counter
import urllib.request, shutil, pathlib, hashlib, json
from openpyxl import load_workbook
n='Costanzo_et_al_Data_File_4_MCMC_consensus_profiles.xlsx';p=pathlib.Path('/tmp/'+n)
u=urllib.request.Request('https://zenodo.org/records/4560674/files/'+n+'?download=1',headers={'User-Agent':'Yeast-Exposed-Encoding-Calibration/1.0'})
with urllib.request.urlopen(u,timeout=180) as inp,p.open('wb') as out:shutil.copyfileobj(inp,out)
assert hashlib.sha256(p.read_bytes()).hexdigest()=='6c0fecfa1b2443a8770b6f4764f1c751318cbd0b2406e06e2a4e9655a737bdbe'
w=load_workbook(p,read_only=True,data_only=True);r=w['Novels_match_v_MCMC_v_Saturated'].iter_rows(values_only=True);next(r)
cs=[Counter(),Counter(),Counter()];sample=[]
for idx,row in enumerate(r):
 for i in range(3):
  v=row[i+5];cs[i][str(type(v).__name__)+'|'+str(v)[:80]]+=1
 if idx<9:sample.append({'genes':[str(x) for x in row[:3]],'raw_reference_values':[{'type':type(v).__name__,'value':repr(v)} for v in row[5:8]]})
w.close()
report={'source_sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'columns':['matched','mcmc','saturated'],
        'value_frequencies_top20':[[[k,n] for k,n in c.most_common(20)] for c in cs],'sample_rows':sample,
        'scope':'EXPOSED CALIBRATION ONLY; V1 readout remains UNKNOWN and not promoted'}
print('CONSENSUS_SOURCE_ENCODING_CALIBRATION='+json.dumps(report,sort_keys=True),flush=True)
