#!/usr/bin/env python3
"""Independent check of source-native MCMC novelty interpretation: no original code imported.

Input mutation is in-memory on the independently parsed, source-hashed Data4 table,
NOT a modification of the publisher's original workbook and NOT a new biological result.
"""
import argparse,collections,hashlib,json,pathlib,shutil,urllib.request
from openpyxl import load_workbook
SRC='Costanzo_et_al_Data_File_4_MCMC_consensus_profiles.xlsx'
SRC_SHA='6c0fecfa1b2443a8770b6f4764f1c751318cbd0b2406e06e2a4e9655a737bdbe'
REPORT_SHA='24d965fa0dabd223ac60af7520e19881b6d5db80b6c0897df9c8cc9d014195b6'
EXPECTED_HEADERS=['Query ORF','Array ORF','Condition','Query Gene','Array Gene','Matched reference','MCMC consensus reference','Saturated reference']
K=('YBR254C_TSQ1398','YDR083W_DMA882','cycloheximide')

def h(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def canon(s):
 x=''.join(c for c in str(s or '').casefold() if c.isalnum())
 return {'boretzeomib':'bortezomib','concanmycina':'concanamycina'}.get(x,x)
def lbl(x):
 if x is None or (isinstance(x,str) and not x.strip()):return 'NOT_LABELED'
 if x=='Positive Novel':return 'POSITIVE_NOVEL'
 if x=='Negative Novel':return 'NEGATIVE_NOVEL'
 raise ValueError('UNDECLARED_SOURCE_LABEL '+repr(x))

def main():
 a=argparse.ArgumentParser();a.add_argument('--v2-report',type=pathlib.Path,required=True);a.add_argument('--out',type=pathlib.Path,required=True);a.add_argument('--cache',type=pathlib.Path,default=pathlib.Path('/tmp/2021-v2-independent-check'))
 v=a.parse_args();v.out.mkdir(parents=True,exist_ok=True);v.cache.mkdir(parents=True,exist_ok=True)
 assert h(v.v2_report)==REPORT_SHA,'ORIGINAL_V2_REPORT_REPLACED'
 report=json.loads(v.v2_report.read_bytes())
 p=v.cache/SRC
 if not p.exists():
  req=urllib.request.Request('https://zenodo.org/records/4560674/files/'+SRC+'?download=1',headers={'User-Agent':'Independent-Biology-Author-Novelty-Check/1.0'})
  with urllib.request.urlopen(req,timeout=180) as r,p.open('wb') as f:shutil.copyfileobj(r,f)
 assert h(p)==SRC_SHA,'ORIGINAL_PUBLISHER_WORKBOOK_REPLACED'
 w=load_workbook(p,read_only=True,data_only=True)
 sh=w['Novels_match_v_MCMC_v_Saturated']; it=sh.iter_rows(values_only=True)
 assert list(next(it))==EXPECTED_HEADERS
 author_records={};counter=collections.Counter()
 for row in it:
  q=str(row[0] or '').strip().upper();b=str(row[1] or '').strip().upper();c=canon(row[2]);key=(q,b,c)
  assert key not in author_records,'DUPLICATE_AUTHORITY_ROW '+str(key)
  lab=tuple(lbl(x) for x in row[5:8]);author_records[key]=lab
  for i,name in enumerate(('matched','mcmc','saturated')):
   if lab[i]!='NOT_LABELED':counter[name]+=1
  counter['source_rows']+=1
 w.close()
 assert counter['source_rows']==3288
 assert counter['matched']==3288 and counter['mcmc']==2682 and counter['saturated']==1769
 assert K in author_records,'NAMED_SELECTED_CELL_MISSING'
 original=author_records[K]
 assert original==('POSITIVE_NOVEL','POSITIVE_NOVEL','NOT_LABELED')
 original_v2=report['protected_selected_event_results']
 assert len(original_v2)==4
 seen={}
 for row in original_v2:
  key=(row['query_allele'],row['array_allele'],row['condition'])
  source=author_records.get(key)
  if source is None:
   assert row['status']=='UNKNOWN_ABSENT_FROM_NOVELS_TABLE'
   seen[key]='UNKNOWN_ABSENT'
  else:
   assert row['status']=='SOURCE_TYPED_ROW_MATCH'
   assert [row[k] for k in ('matched_novelty_label','mcmc_novelty_label','saturated_novelty_label')]==list(source)
   seen[key]='SOURCE_MATCH'
 assert collections.Counter(seen.values())=={'UNKNOWN_ABSENT':3,'SOURCE_MATCH':1}
 # Simulate a prospective *authority input* mutation on one independently parsed
 # record only: change its MCMC classification to blank. This does not change
 # the biological differential effect or the matched/saturated attributions.
 copy=dict(author_records)
 copy[K]=(original[0],'NOT_LABELED',original[2]);assert copy[K]!=original
 assert copy[K][0]==original[0] and copy[K][2]==original[2]
 changed=[key for key in author_records if author_records[key]!=copy[key]]
 assert changed==[K]
 before={r:sum(v[i]!='NOT_LABELED' for v in author_records.values()) for i,r in enumerate(('matched','mcmc','saturated'))}
 after={r:sum(v[i]!='NOT_LABELED' for v in copy.values()) for i,r in enumerate(('matched','mcmc','saturated'))}
 assert before=={'matched':3288,'mcmc':2682,'saturated':1769}
 assert after=={'matched':3288,'mcmc':2681,'saturated':1769}
 copy[K]=original
 assert copy==author_records,'RESTORATION_DID_NOT_REQUALIFY_SOURCE_VIEW'
 res={'status':'INDEPENDENT_SOURCE_REPLAY_AND_SINGLE_REFERENCE_CHANGE_PASS','source_sha256':SRC_SHA,'v2_report_sha256':REPORT_SHA,
       'native_author_record_count':counter['source_rows'],'reference_novelty_label_counts':before,
       'protected_selected_source_events':{'matched_exact':1,'unknown_not_in_novelty_table':3},
       'selected_exact_source_row':{'identity':K,'original_labels':original},
       'counterfactual_in_memory_authority_value_update':{'changed_records':1,'before':before,'after':after,'matched_and_saturated_unchanged':True,'exact_restore':True},
       'limits':['Already opened original V3 source outcomes; this is posthoc independent checking, not new experiment',
                 'In-memory derived source view change, NOT edit of published workbook or true revocation of publisher statement',
                 'No claim of full CLC / Flash kernel or biological method superiority']}
 (v.out/'independent_reference_audit.json').write_text(json.dumps(res,indent=2,sort_keys=True)+'\n')
 print('INDEPENDENT_BIOLOGY_REFERENCE_AUDIT='+json.dumps(res,sort_keys=True),flush=True)
 print('SOURCE_NATIVE_NOVELTY_AND_CAUSAL_CHANGE_CONTROL_PASS',flush=True)
 if not v.cache.joinpath('KEEP').exists():p.unlink()
if __name__=='__main__':main()