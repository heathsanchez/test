#!/usr/bin/env python3
"""Only source workbook schema/metadata; absolutely no consensus outcome values read."""
import hashlib, json, pathlib, urllib.request, shutil
from openpyxl import load_workbook
name='Costanzo_et_al_Data_File_4_MCMC_consensus_profiles.xlsx'
url='https://zenodo.org/records/4560674/files/'+name+'?download=1'
dest=pathlib.Path('/tmp/'+name)
request=urllib.request.Request(url,headers={'User-Agent':'Biology-Consensus-Schema-Only/1.0'})
with urllib.request.urlopen(request,timeout=180) as src,dest.open('wb') as dst:
    shutil.copyfileobj(src,dst)
raw=dest.read_bytes()
digest=hashlib.md5(raw).hexdigest()
assert digest=='686f8b88507826d3cb7117e80274be92','WRONG_CONSENSUS_SOURCE'
wb=load_workbook(dest,read_only=True,data_only=False)
metadata=[]
for sheet in wb:
    header=next(sheet.iter_rows(min_row=1,max_row=1,values_only=True),())
    metadata.append({'sheet':sheet.title,'max_row':sheet.max_row,'max_column':sheet.max_column,
                     'first_row_30':[str(x)[:80] if x is not None else '' for x in header[:30]]})
wb.close()
report={'source':name,'md5':digest,'sha256':hashlib.sha256(raw).hexdigest(),
        'bytes':len(raw),'metadata_only':True,'sheets':metadata}
print('CONSENSUS_2021_METADATA_ONLY='+json.dumps(report,sort_keys=True),flush=True)
