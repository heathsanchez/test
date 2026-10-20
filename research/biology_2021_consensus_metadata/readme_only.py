#!/usr/bin/env python3
"""Only source-authored README text; no consensus rows accessed."""
import json, pathlib, urllib.request, shutil, hashlib
from openpyxl import load_workbook
name='Costanzo_et_al_Data_File_4_MCMC_consensus_profiles.xlsx'
p=pathlib.Path('/tmp/'+name)
with urllib.request.urlopen(urllib.request.Request('https://zenodo.org/records/4560674/files/'+name+'?download=1',headers={'User-Agent':'Biology-Consensus-Readme-Only/1.0'}),timeout=180) as inp,p.open('wb') as out:shutil.copyfileobj(inp,out)
assert hashlib.sha256(p.read_bytes()).hexdigest()=='6c0fecfa1b2443a8770b6f4764f1c751318cbd0b2406e06e2a4e9655a737bdbe'
wb=load_workbook(p,read_only=True,data_only=True);ws=wb['Readme']
lines=[[str(v)[:500] if v is not None else '' for v in row[:2]] for row in ws.iter_rows(min_row=1,max_row=ws.max_row,values_only=True)]
wb.close()
print('CONSENSUS_SOURCE_AUTHORED_README='+json.dumps(lines),flush=True)
