#!/usr/bin/env python3
"""Reads ONLY environmental condition labels from official 2021 conditions workbook."""
from urllib.request import urlopen,Request
from openpyxl import load_workbook
import hashlib,pathlib,os,json
name="Costanzo_et_al_Data_File_1_Conditions_Strains_Fitness.xlsx"
url="https://zenodo.org/records/4560674/files/"+name+"?download=1"
p=pathlib.Path("/tmp/"+name)
with urlopen(Request(url,headers={"User-Agent":"metalogic-2021-condition-labels"}),timeout=120) as r,p.open("wb") as f:
 import shutil;shutil.copyfileobj(r,f)
assert hashlib.md5(p.read_bytes()).hexdigest()=="feec6ee93cbdc6c29d1989b05bb15390"
wb=load_workbook(p,read_only=True,data_only=True)
ws=wb["Conditions"]
vals=[str(row[0]).strip() for row in ws.iter_rows(min_row=2,values_only=True) if row[0] is not None and str(row[0]).strip()]
print("EXTERNAL_2021_CONDITION_LABELS="+json.dumps(vals),flush=True)
wb.close()
