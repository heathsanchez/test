#!/usr/bin/env python3
"""2021 Costanzo environmental GI source metadata preflight. No numeric data outcomes accessed."""
import hashlib, json, os, pathlib, urllib.request, zipfile
import openpyxl
FILES = {
"conditions":("Costanzo_et_al_Data_File_1_Conditions_Strains_Fitness.xlsx","feec6ee93cbdc6c29d1989b05bb15390"),
"interactions":("Costanzo_et_al_Data_File_3_Raw_interaction_dataset.xlsx","f81a8123a8113b795a00e8e154fa072d"),
}
BASE="https://zenodo.org/records/4560674/files/"
p=pathlib.Path("out/environment_metadata");p.mkdir(parents=True,exist_ok=True)
out={"source_doi":"10.5061/dryad.3r2280gfd","archive":"Zenodo record 4560674","disclosure_boundary":"Only workbook sheet identities, first header row, declared dimensions; NO row outcomes or data statistics","files":{}}
for role,(fn,md5) in FILES.items():
 dest=p/fn
 req=urllib.request.Request(BASE+fn+"?download=1",headers={"User-Agent":"Metalogic-Biology-Metadata-Audit/1.0"})
 with urllib.request.urlopen(req,timeout=300) as inp,dest.open("wb") as f:
  import shutil;shutil.copyfileobj(inp,f)
 digest=hashlib.md5(dest.read_bytes()).hexdigest()
 assert digest==md5, f"AUTHORITY_HASH_MISMATCH {role} {digest}"
 rec={"name":fn,"md5":digest,"sha256":hashlib.sha256(dest.read_bytes()).hexdigest(),"size":dest.stat().st_size,"sheets":[]}
 wb=openpyxl.load_workbook(dest,read_only=True,data_only=False)
 for sheet in wb:
  it=sheet.iter_rows(min_row=1,max_row=1,values_only=True)
  header=next(it,None)
  values=[str(v)[:120] if v is not None else "" for v in header] if header is not None else []
  rec["sheets"].append({"name":sheet.title,"max_row":sheet.max_row,"max_column":sheet.max_column,
                        "first_row_first_30_headers":values[:30],"nonempty_header_count":sum(bool(v) for v in values)})
 wb.close()
 out["files"][role]=rec
 dest.unlink()
print("2021_ENVIRONMENT_METADATA_ONLY="+json.dumps(out,sort_keys=True),flush=True)
(p/"source_metadata.json").write_text(json.dumps(out,indent=2,sort_keys=True)+"\n")
