#!/usr/bin/env python3
"""Independent native 2021 environmental GI evaluator: reads raw source only AFTER predictions sealed.
Does NOT import predictor or invoke any of its functions. Two study replicates and
source-authored P values are retained; P values are NOT tests of cross-gene contrasts.
"""
import argparse
from collections import defaultdict,Counter
from decimal import Decimal,InvalidOperation
import hashlib
import json
from pathlib import Path
import numpy as np
from openpyxl import load_workbook
import urllib.request, shutil

NAME='Costanzo_et_al_Data_File_3_Raw_interaction_dataset.xlsx'
MD5='f81a8123a8113b795a00e8e154fa072d'
SHA='b4e242087e554df1adc648aefb890144d397a1eb303a55509a90ae08e912c488'
D=Decimal


def source(cache):
 cache.mkdir(parents=True,exist_ok=True);p=cache/NAME
 if not p.exists():
  with urllib.request.urlopen(urllib.request.Request('https://zenodo.org/records/4560674/files/'+NAME+'?download=1',headers={'User-Agent':'Metalogic-Independent-2021-Scorer'}),timeout=360) as r,p.open('wb') as f:shutil.copyfileobj(r,f)
 h=hashlib.sha256();m=hashlib.md5()
 with p.open('rb') as f:
  for b in iter(lambda:f.read(1024*1024),b''):h.update(b);m.update(b)
 if h.hexdigest()!=SHA or m.hexdigest()!=MD5:raise RuntimeError('SOURCE_ERROR')
 return p


def dec(v):
 if v is None or v=='':return None
 try:
  a=D(str(v))
  return a if a.is_finite() else None
 except (InvalidOperation,TypeError,ValueError):return None


def classify(rec):
 if rec is None:return None,None
 e,p,r1,r2=rec
 if e is None or p is None:return None,None
 main=bool(p<D('.05') and abs(e)>D('.08'))
 strict=bool(p<D('.005') and abs(e)>D('.08') and r1 is not None and r2 is not None
             and r1*r2>0 and e*r1>0 and e*r2>0)
 return main,strict


def bootstrap_query_cluster(eff,qs,B=4000):
 # Resample independent query-gene units to keep same-query pair dependence clustered.
 kinds=sorted(set(qs)); rng=np.random.default_rng(20261020)
 if len(kinds)<4:return [None,None]
 by={q:[eff[i] for i in range(len(qs)) if qs[i]==q] for q in kinds}
 samples=[]
 for _ in range(B):
  keys=rng.choice(kinds,len(kinds),replace=True)
  arr=[v for k in keys for v in by[k]]
  samples.append(float(np.mean(arr)))
 return [round(float(x),8) for x in np.quantile(samples,[.025,.975])]


def main():
 p=argparse.ArgumentParser();p.add_argument('--allowed',type=Path,required=True);p.add_argument('--sealed',type=Path,required=True);p.add_argument('--out',type=Path,required=True);p.add_argument('--cache',type=Path,default=Path('/tmp/environment2021-evaluator'))
 a=p.parse_args();a.out.mkdir(parents=True,exist_ok=True)
 ab=a.allowed.read_bytes();pb=a.sealed.read_bytes()
 allowed=json.loads(ab);sealed=json.loads(pb)
 ah=hashlib.sha256(ab).hexdigest();ph=hashlib.sha256(pb).hexdigest()
 assert sealed['allowed_sha256']==ah and sealed['source_sha256']==SHA
 assert sealed['protocol']==allowed['protocol']=='2021_environment_v3'
 targets=allowed['target'];decisions=sealed['decisions'];assert len(targets)==len(decisions)
 if [x['id'] for x in targets]!=[x['id'] for x in decisions]:raise RuntimeError('TARGET_ID_MISMATCH')
 known=allowed['known_conditions'];candidates=allowed['candidate_conditions'];contexts=set(known+candidates)
 target_ids={z['id'] for z in targets};native=defaultdict(list)
 path=source(a.cache)
 wb=load_workbook(path,read_only=True,data_only=True);ws=wb['Diagnostic array_complete']
 it=ws.iter_rows(values_only=True);headers=tuple(next(it))
 if len(headers)!=25 or headers[17]!='mean_differential_epsilon' or headers[19]!='differential_p_value':raise RuntimeError('SCHEMA_MISMATCH')
 for row in it:
  q=str(row[0] or '').strip().upper();arr=str(row[1] or '').strip().upper();c=' '.join(str(row[2] or '').strip().split()).lower()
  c={'boretzeomib':'bortezomib','concanmycin a':'concanamycin a'}.get(c,c)
  key=q+'|'+arr
  if key not in target_ids or c not in contexts:continue
  e,pv,r1,r2=dec(row[17]),dec(row[19]),dec(row[15]),dec(row[16])
  ref=dec(row[12])
  native[(key,c)].append(((e,pv,r1,r2),ref))
 wb.close()
 # Original observed features are separately reconstructed from native original bytes.
 for t in targets:
  observed=[native.get((t['id'],c),[]) for c in known]
  if any(len(x)!=1 for x in observed):raise RuntimeError('INDEPENDENT_KNOWN_CELL_NOT_UNIQUE')
  for idx,x in enumerate(observed):
   v=x[0][0]
   if v[0] is None or abs(float(v[0])-t['x'][idx])>1e-9 or classify(v)[0]:
    raise RuntimeError('TARGET_INITIAL_PANEL_REPLAY_MISMATCH '+t['id'])
  rs=[v[0][1] for v in observed]
  if any(x is None for x in rs):raise RuntimeError('REFERENCE_MISSING')
  ref=sum(rs)/D(4)
  if abs(float(ref)-t['x'][4])>1e-9:raise RuntimeError('SOURCE_FEATURE_REPLAY_MISMATCH')
 method_names=sealed['methods'];totals={m:{'attempts':0,'observed':0,'success':0,'strict_success':0,'missing':0} for m in method_names}
 successes={m:[] for m in method_names};traces=[];oracle=0;oracle_strict=0; qids=[]
 for t,d in zip(targets,decisions):
  qids.append(t['q']);poss=[];strictposs=[];obs={}
  for c in candidates:
   matches=native.get((t['id'],c),[])
   if len(matches)!=1:obs[c]=(None,None,None,None)
   else:obs[c]=matches[0][0]
   result,strict=classify(obs[c]);poss.append(result is True);strictposs.append(strict is True)
  oracle+=int(any(poss));oracle_strict+=int(any(strictposs))
  tr={'id':t['id'],'q':t['q'],'a':t['a'],'chosen':{},'oracle_some_candidate':any(poss),'oracle_strict_some_candidate':any(strictposs)}
  for m in method_names:
   ctx=d['choices'][m]
   if ctx not in candidates:raise RuntimeError('UNAUTHORIZED_OUTCOME_CHANNEL')
   result,strict=classify(obs[ctx]);counter=totals[m];counter['attempts']+=1
   counter['observed']+=int(result is not None);counter['missing']+=int(result is None)
   counter['success']+=int(result is True);counter['strict_success']+=int(strict is True)
   successes[m].append(int(result is True))
   e,p,r1,r2=obs[ctx]
   tr['chosen'][m]={'condition':ctx,'measured':result is not None,'source_effect':str(e) if e is not None else None,
      'source_p':str(p) if p is not None else None,'signed_replicates':[str(x) if x is not None else None for x in [r1,r2]],'hit':result,'strict_hit':strict}
  traces.append(tr)
 p1='contrast_exploration';p2='knn_hit';deltas=np.array(successes[p1],dtype=int)-np.array(successes[p2],dtype=int)
 for m in totals:
  totals[m]['rate_all_attempts']=totals[m]['success']/len(targets)
  totals[m]['strict_rate_all_attempts']=totals[m]['strict_success']/len(targets)
 result={'status':'REAL_ENVIRONMENTAL_GI_HOLDOUT_EVALUATED','source_study':'Costanzo et al 2021 Environmental robustness of the global yeast genetic interaction network',
  'source_doi':'10.5061/dryad.3r2280gfd','source_sha256':SHA,'allowed_sha256':ah,
  'sealed_prediction_sha256':ph,'n_units':len(targets),'distinct_target_queries':len(set(qids)),
  'expected_no_train_target_gene_overlap':allowed['separation']['no_overlap'],
  'source_significance_rule':'p<0.05 and abs(mean differential epsilon)>0.08',
  'strict_secondary_rule':'p<0.005 and abs(mean differential epsilon)>0.08 and both replicate signs agree',
  'methods':totals,'oracle_measurable_targets':oracle,'oracle_strict_targets':oracle_strict,
  'primary_method_difference':float(deltas.mean()),'query_cluster_bootstrap_95':bootstrap_query_cluster(deltas,qids),
  'no_claims':['new biological experiment','statistical independence of genotypes tested under same query','corrected multiple-testing familywise significance for all source hypotheses','novel law','evidence that human-authored selector is uniquely optimal']}
 (a.out/'outcome_report.json').write_text(json.dumps(result,indent=2,sort_keys=True)+'\n')
 (a.out/'per_target_receipt.json').write_text(json.dumps(traces,indent=2,sort_keys=True)+'\n')
 print('ENVIRONMENT_2021_NATIVE_OUTCOME_RESULT='+json.dumps(result,sort_keys=True),flush=True)
 print('ENVIRONMENT_2021_INDEPENDENT_SOURCE_EVALUATION_PASS',flush=True)

if __name__=='__main__':main()
