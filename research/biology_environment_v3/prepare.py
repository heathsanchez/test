#!/usr/bin/env python3
"""Prepare a strictly limited feature-only evidence view from verified 2021 environmental GI data.

This source steward sees the physical workbook, but deliberately DOES NOT READ target
candidate-outcome columns or emit any target candidate scores, P-values or missingness.
The machine predicting interventions runs in a separate GitHub Actions job receiving
ONLY the produced allowlisted JSON. Histories/unknowns are not silently excluded.
"""
import argparse
from collections import defaultdict, Counter
import gzip
import hashlib
import json
import math
from pathlib import Path
import re
import shutil
from urllib.request import Request, urlopen
from openpyxl import load_workbook

DOI = '10.5061/dryad.3r2280gfd'
BASE = 'https://zenodo.org/records/4560674/files/'
CONDITIONS_FILE = 'Costanzo_et_al_Data_File_1_Conditions_Strains_Fitness.xlsx'
RESULTS_FILE = 'Costanzo_et_al_Data_File_3_Raw_interaction_dataset.xlsx'
EXPECTED = {
 CONDITIONS_FILE: ('feec6ee93cbdc6c29d1989b05bb15390','02d97bf2c57397c781e4c7792430f89938c92da823df58836b53185cb62b4655'),
 RESULTS_FILE: ('f81a8123a8113b795a00e8e154fa072d','b4e242087e554df1adc648aefb890144d397a1eb303a55509a90ae08e912c488')
}
HEADER = ('query_orf','array_orf','Condition','query_common','array_common',
'rep1_condition_epsilon','rep2_condition_epsilon','mean_condition_epsilon','sd_condition_epsilon','condition_p_value',
'rep1_reference_epsilon','rep2_reference_epsilon','mean_reference_epsilon','sd_reference_epsilon','reference_p_value',
'rep1_differential_epsilon','rep2_differential_epsilon','mean_differential_epsilon','sd_differential_epsilon',
'differential_p_value','Differential Interaction Sign','Differential interaction class','Query-Array',
'epsilon_Costanzo et al 2016','pvalue_Costanzo et al 2016')
ORF = re.compile(r'^Y[A-P][LR]\d{3}[CW](?:-[A-Z])?$', re.I)
SEED = 'biology-environment-v3-20261020'
MAX_TARGETS = 200


def digest(path):
 h=hashlib.md5();s=hashlib.sha256()
 with open(path,'rb') as f:
  for x in iter(lambda:f.read(1024*1024),b''):h.update(x);s.update(x)
 return h.hexdigest(),s.hexdigest()


def source(cache):
 cache.mkdir(parents=True,exist_ok=True)
 files={}
 for name,want in EXPECTED.items():
  dest=cache/name
  if not dest.exists():
   with urlopen(Request(BASE+name+'?download=1',headers={'User-Agent':'Metalogic-source-steward-v3/1.0'}),timeout=360) as u, dest.open('wb') as f:
    shutil.copyfileobj(u,f)
  got=digest(dest)
  if got!=want: raise RuntimeError('SOURCE_HASH_MISMATCH '+name+' '+str(got))
  files[name]=str(dest)
 return files


def norm(s):
 key=' '.join(str(s or '').strip().split()).lower()
 # Official 2021 workbook spellings differ from interaction-record spellings.
 return {'boretzeomib':'bortezomib','concanmycin a':'concanamycin a'}.get(key,key)


def validfloat(x):
 if x is None or isinstance(x,bool):return None
 try:
  v=float(x)
  return v if math.isfinite(v) else None
 except (TypeError,ValueError):return None


def side(orf):
 v=int.from_bytes(hashlib.sha256((SEED+'|'+orf.upper()).encode()).digest()[:8],'big')
 return 'calibration' if v / 2**64 <0.70 else 'target'


def hit(row):
 # Publication's intermediate per-interaction confidence: P<.05 and |differential epsilon|>.08
 # NOT a multiple-testing-adjusted claim; strict secondary is evaluated separately.
 e=validfloat(row[17]);p=validfloat(row[19])
 if e is None or p is None:return None
 return bool(p<.05 and abs(e)>.08)


def known_conditions(conditions_file):
 wb=load_workbook(conditions_file,read_only=True,data_only=True)
 ws=wb['Conditions'];ls=[]
 for row in ws.iter_rows(min_row=2,values_only=True):
  s=norm(row[0])
  if s:ls.append(s)
 wb.close()
 ls=sorted(set(x for x in ls if x!='reference condition + solvent'))
 if len(ls)!=14:raise RuntimeError('SCHEMA_CONDITION_COUNT_NOT_14 '+repr(ls))
 return ls[:4],ls[4:],ls


def save_json(path,data):
 b=(json.dumps(data,sort_keys=True,separators=(',',':'),allow_nan=False)+'\n').encode()
 path.parent.mkdir(parents=True,exist_ok=True);path.write_bytes(b)
 return hashlib.sha256(b).hexdigest()


def build(fixtures,output,fixture_only=False):
 known,candidates,all_conditions=known_conditions(fixtures[CONDITIONS_FILE])
 record={};count=Counter();duplicate_keys=Counter();obs_conditions=set()
 wb=load_workbook(fixtures[RESULTS_FILE],read_only=True,data_only=True)
 ws=wb['Diagnostic array_complete']
 rows=ws.iter_rows(values_only=True);header=tuple(next(rows))
 if header!=HEADER:raise RuntimeError('SCHEMA_25_COLUMN_IDENTITY_MISMATCH '+repr(header))
 if not fixture_only and ws.max_row !=363202:raise RuntimeError('SCHEMA_ROW_COUNT_MISMATCH '+str(ws.max_row))
 for row in rows:
  q,a,c=map(norm,row[:3]);q=q.upper();a=a.upper()
  if c not in all_conditions:
   if c:raise RuntimeError('UNRECOGNISED_CONDITION '+c)
   continue
  if not ORF.fullmatch(q) or not ORF.fullmatch(a) or q==a:
   count['unusable_gene_identity']+=1;continue
  qs=side(q);as_=side(a)
  if qs!=as_:
   count['cross_side_removed']+=1;continue
  # CRITICAL: target candidate numeric results are never accessed in Python.
  if qs=='target' and c not in known:
   # No numeric read, no missingness/coverage counters and no result metadata.
   continue
  ident=q+'|'+a
  obs_conditions.add(c)
  item=record.setdefault((qs,ident),{'q':q,'a':a,'seen':{}})
  if c in item['seen']:
   duplicate_keys[(qs,ident,c)]+=1;continue
  # Only source steward's calibration data + target known-panel values are interpreted.
  e=validfloat(row[17]);p=validfloat(row[19]);ref=validfloat(row[12]);rep1=validfloat(row[15]);rep2=validfloat(row[16])
  item['seen'][c]={'epsilon':e,'p':p,'reference':ref,'replicate1':rep1,'replicate2':rep2,
                   'hit':hit(row)}
  count['allowed_observations_interpreted']+=1
 wb.close()
 if set(all_conditions)!=obs_conditions:raise RuntimeError('NOT_ALL_14_CONDITIONS_PRESENT '+str(sorted(obs_conditions)))
 # Duplicate source keys are counted as UNKNOWN outcomes, never overwritten by last value.
 for (role,ident,condition),n in duplicate_keys.items():
  record[(role,ident)]['seen'][condition]=None
 train=[];target=[]
 for (role,ident),item in record.items():
  seen=item['seen']
  observed=[seen.get(c) for c in known]
  if any(t is None or t['epsilon'] is None or t['p'] is None or t['reference'] is None for t in observed):
   count['missing_known_panel_or_duplicate']+=1;continue
  x=[t['epsilon'] for t in observed]+[sum(t['reference'] for t in observed)/4.0]
  if any(not math.isfinite(v) for v in x):raise AssertionError('Nonfinite feature')
  if role=='calibration':
   y=[];effect=[]
   for c in candidates:
    v=seen.get(c)
    y.append(None if v is None else v['hit'])
    effect.append(None if v is None else v['epsilon'])
   if all(yv is None for yv in y):
    count['calibration_no_candidate_outcomes']+=1;continue
   train.append({'id':ident,'q':item['q'],'a':item['a'],'x':x,'y':y,'effect':effect})
  else:
   if any(t['hit'] for t in observed):
    count['target_already_responding_in_initial_panel']+=1;continue
   # Target record intentionally includes NO candidate result/status/outcome flags.
   target.append({'id':ident,'q':item['q'],'a':item['a'],'x':x})
 train.sort(key=lambda t:t['id'])
 target.sort(key=lambda t:hashlib.sha256((SEED+'|choice|'+t['id']).encode()).hexdigest())
 target=target[:MAX_TARGETS]
 if len(train)<1000 or len(target)<40:
  raise RuntimeError('INAPPLICABLE_TOO_FEW_INDEPENDENT_TRAIN_OR_TEST_GENES '
                     +json.dumps({'train':len(train),'target':len(target)}))
 cg={z for i in train for z in (i['q'],i['a'])}
 tg={z for i in target for z in (i['q'],i['a'])}
 assert not(cg & tg),'TRAIN_TEST_GENE_LEAKAGE'
 # No protected target candidate score, P value, missingness or replicate is materialized.
 allowed={'protocol':'2021_environment_v3','source_doi':DOI,
 'source_sha256':EXPECTED[RESULTS_FILE][1],
 'source_conditions_sha256':EXPECTED[CONDITIONS_FILE][1],
 'selection_seed':SEED,'known_conditions':known,'candidate_conditions':candidates,
 'x_features':'four source-mapped known differential epsilon plus 2021 untreated reference mean',
 'classification':'P<0.05 and absolute mean differential epsilon>0.08',
 'training':train,'target':target,
 'separation':{'train_gene_count':len(cg),'target_gene_count':len(tg),'no_overlap':True},
 'source_steward_counts':dict(count),'duplicates_observed_in_allowed_records':len(duplicate_keys)}
 h=save_json(output/'allowed.json',allowed)
 print('BIOLOGY_ENVIRONMENT_V3_ALLOWED_SHA256='+h,flush=True)
 print('BIOLOGY_ENVIRONMENT_V3_ALLOWED_COUNTS='+json.dumps({'train':len(train),'target':len(target),'train_genes':len(cg),'target_genes':len(tg),'known':known,'candidates':candidates,'source_sha256':allowed['source_sha256'],'guard':'no target candidate outcome ever interpreted or serialized'},sort_keys=True),flush=True)
 return allowed


def main():
 p=argparse.ArgumentParser();p.add_argument('--output',type=Path,required=True);p.add_argument('--cache',type=Path,default=Path('/tmp/environment2021cache'));p.add_argument('--fixture-dir',type=Path)
 args=p.parse_args();args.output.mkdir(parents=True,exist_ok=True)
 f={name:str(args.fixture_dir/name) for name in EXPECTED} if args.fixture_dir else source(args.cache)
 build(f,args.output,fixture_only=bool(args.fixture_dir))

if __name__=='__main__':main()
