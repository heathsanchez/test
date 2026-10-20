#!/usr/bin/env python3
"""Adversarial fixture: protected target candidate outcomes are physically changed,
including P-values, signs and missingness; predictor's bytes and choices must be invariant.
All source data are synthetic and must NEVER be reported as biological results.
"""
import hashlib
import json
from pathlib import Path
import tempfile
from openpyxl import Workbook
import prepare
import predict
import evaluate


def names(side,n,forbidden):
 result=[]
 for chrid in range(16):
  for i in range(1,200):
   for arm in 'LR':
    g=f'Y{chr(65+chrid)}{arm}{i:03d}W'
    if g not in forbidden and prepare.side(g)==side:
     forbidden.add(g);result.append(g)
     if len(result)==n:return result
 raise AssertionError('not enough synthetic names')


def gen(dest,poison=False):
 dest.mkdir(parents=True,exist_ok=True)
 obs=['Reference condition + solvent','Actinomycin D','Benomyl','Boretzeomib','Caspofungin','Concanmycin A','Cycloheximide','Fluconazole','Galactose','Geldanamycin','MMS','Monensin','Rapamycin','Sorbitol','Tunicamycin']
 k=Workbook();w=k.active;w.title='Conditions';w.append(['Condition','Concentration'])
 for nm in obs:w.append([nm,'unneeded'])
 k.save(dest/prepare.CONDITIONS_FILE)
 g=Workbook(write_only=True);ws=g.create_sheet('Diagnostic array_complete');ws.append(prepare.HEADER)
 forbidden=set()
 qs={s:names(s,18 if s=='calibration' else 8,forbidden) for s in ('calibration','target')}
 arrays={s:names(s,80 if s=='calibration' else 52,forbidden) for s in ('calibration','target')}
 idx=0
 for role in ('calibration','target'):
  for q in qs[role]:
   for a in arrays[role]:
    idx+=1
    for ci,c in enumerate(obs[1:]):
     known=c.casefold() in ['actinomycin d','benomyl','boretzeomib','caspofungin']
     big=(idx*7+ci*13)%8==0 and not known
     v=round(0.18 if big else 0.014+(((idx*3+ci)%9)-4)*0.002,4)
     p=0.008 if big else 0.45
     # Deterministic adversarial corruption can turn a measured hit into missing
     # or reverse its sign. None of these protected target values may affect selection.
     if role=='target' and not known and poison:
      if (idx+ci)%2==0:v=None;p=None
      else:v=-0.4;p=0.0000001
     row=[q,a,c,'q','a',0,0,0,0,None,0,0,0.02,0.001,0.5,
          v,v,v,0.002 if v is not None else None,p,'','','',0.04,0.4]
     ws.append(row)
 g.save(dest/prepare.RESULTS_FILE)
 return qs,arrays


def run():
 with tempfile.TemporaryDirectory() as d:
  base=Path(d)
  fa=base/'fixtureA';fb=base/'fixtureB'
  qs,arr=gen(fa,False);gen(fb,True)
  da=base/'allowA';db=base/'allowB'
  a=prepare.build({n:str(fa/n) for n in prepare.EXPECTED},da,fixture_only=True)
  b=prepare.build({n:str(fb/n) for n in prepare.EXPECTED},db,fixture_only=True)
  bytea=(da/'allowed.json').read_bytes();byteb=(db/'allowed.json').read_bytes()
  assert bytea==byteb, 'SOURCE_STEWARD_LEAKED_PROTECTED_CANDIDATES'
  assert len(a['training'])>=1000 and len(a['target'])>=40
  assert all(set(t)=={'id','q','a','x'} for t in a['target'])
  assert all(all(k not in t for k in ('y','effect','candidate_scores','candidate_missingness')) for t in a['target'])
  cg={g for t in a['training'] for g in (t['q'],t['a'])}
  tg={g for t in a['target'] for g in (t['q'],t['a'])}
  assert not (cg & tg)
  assert len(a['known_conditions'])==4 and len(a['candidate_conditions'])==10
  choices1=predict.rank_contexts(a);choices2=predict.rank_contexts(b)
  assert choices1==choices2,'PREDICTOR_CHOICES_CHANGED_BY_HIDDEN_OUTCOMES'
  for z in choices1:
   assert len(z['choices'])==len(predict.METHODS)
   assert set(z['choices'].values())<=set(a['candidate_conditions'])
  assert evaluate.classify((evaluate.D('.1'),evaluate.D('.01'),evaluate.D('.1'),evaluate.D('.1'))) == (True,False)
  assert evaluate.classify((evaluate.D('.1'),evaluate.D('.001'),evaluate.D('.1'),evaluate.D('.1'))) == (True,True)
  assert evaluate.classify((evaluate.D('.08'),evaluate.D('.0001'),evaluate.D('.1'),evaluate.D('.1'))) == (False,False)
  assert evaluate.classify((evaluate.D('.160'),evaluate.D('.05'),evaluate.D('.1'),evaluate.D('.1'))) == (False,False)
  assert evaluate.classify(None)==(None,None)
  print('BIOLOGY_V3_SYNTHETIC_TARGET_OUTCOME_ABLATION_PASS='+json.dumps({'train':len(a['training']), 'targets':len(a['target']), 'allowed_sha256':hashlib.sha256(bytea).hexdigest(),'no_shared_gene':True,'protected_perturbation':'alter every target candidate response including P values and missingness','n_methods':len(predict.METHODS)},sort_keys=True),flush=True)

if __name__=='__main__':run()
